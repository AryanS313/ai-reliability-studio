from __future__ import annotations

import json
import re
import zipfile
from collections.abc import Sequence
from io import BytesIO
from pathlib import Path
from typing import Any

from src import config
from src.infrastructure import MalwareScanner, OCRPipeline, ScanVerdict
from src.security import validate_upload
from src.versioning import text_hash

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".log",
    ".yaml",
    ".yml",
    ".xml",
    ".pdf",
    ".docx",
    ".rtf",
    ".csv",
    ".tsv",
    ".xlsx",
    ".xls",
    ".json",
    ".jsonl",
    ".html",
    ".htm",
    ".pptx",
}
UPLOAD_TYPES = sorted(extension.lstrip(".") for extension in SUPPORTED_EXTENSIONS)
ZIP_BASED_EXTENSIONS = {".docx", ".xlsx", ".pptx"}


def load_document(
    path: str | Path,
    *,
    malware_scanner: MalwareScanner | None = None,
    ocr_pipeline: OCRPipeline | None = None,
) -> dict[str, Any]:
    file_path = Path(path)
    artifact = extract_document_artifact(
        file_path.name,
        file_path.read_bytes(),
        malware_scanner=malware_scanner,
        ocr_pipeline=ocr_pipeline,
    )
    return {**artifact, "path": str(file_path)}


def load_documents(paths: Sequence[str | Path]) -> list[dict[str, Any]]:
    return [load_document(path) for path in paths]


def load_uploaded_document(
    uploaded,
    *,
    malware_scanner: MalwareScanner | None = None,
    ocr_pipeline: OCRPipeline | None = None,
) -> dict[str, Any]:
    artifact = extract_document_artifact(
        uploaded.name,
        uploaded.getvalue(),
        malware_scanner=malware_scanner,
        ocr_pipeline=ocr_pipeline,
    )
    return {**artifact, "path": artifact["filename"]}


def extract_document_bytes(filename: str, data: bytes) -> str:
    return str(extract_document_artifact(filename, data)["text"])


def extract_document_artifact(
    filename: str,
    data: bytes,
    *,
    malware_scanner: MalwareScanner | None = None,
    ocr_pipeline: OCRPipeline | None = None,
) -> dict[str, Any]:
    if Path(str(filename)).suffix.lower() == ".doc":
        raise ValueError("Legacy .doc files are not supported reliably. Save the file as .docx and upload it again.")
    safe_name = validate_upload(filename, data, allowed_extensions=SUPPORTED_EXTENSIONS)
    suffix = Path(safe_name).suffix.lower()
    warnings: list[str] = []
    scan_metadata: dict[str, Any]
    if malware_scanner is None:
        if config.REQUIRE_MALWARE_SCAN:
            raise ValueError("Production document ingestion requires a configured malware scanner.")
        scan_metadata = {"verdict": "not_requested", "scanner": None, "production_ready": False}
    else:
        scan = malware_scanner.scan(data, safe_name)
        scan_metadata = {
            "verdict": scan.verdict.value,
            "scanner": scan.scanner,
            "signature_version": scan.signature_version,
            "production_ready": malware_scanner.production_ready,
        }
        if scan.verdict == ScanVerdict.MALICIOUS:
            raise ValueError("Document was rejected by malware scanning.")
        if scan.verdict == ScanVerdict.UNAVAILABLE:
            if config.REQUIRE_MALWARE_SCAN:
                raise ValueError("Malware scanning is unavailable; production ingestion is fail-closed.")
            warnings.append(scan.safe_reason or "Malware scanning was unavailable.")
    if suffix in ZIP_BASED_EXTENSIONS:
        _validate_zip_container(data)
    blocks: list[dict[str, Any]] = []
    try:
        if suffix in {".txt", ".md", ".markdown", ".log", ".yaml", ".yml", ".xml"}:
            text = _decode_text(data)
            blocks = _text_blocks(text, safe_name)
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))
            pages = []
            for page_number, page in enumerate(reader.pages, start=1):
                page_text = (page.extract_text() or "").strip()
                if not page_text:
                    warnings.append(
                        f"Page {page_number} contains no searchable text; configure the OCR extension for scanned pages."
                    )
                    continue
                pages.append(f"Page {page_number}\n{page_text}")
                blocks.extend(_text_blocks(page_text, safe_name, page=page_number))
            text = "\n\n".join(pages)
            if not text.strip() and ocr_pipeline is not None:
                ocr = ocr_pipeline.extract(data, filename=safe_name)
                if ocr.available and ocr.text.strip():
                    text = ocr.text.strip()
                    blocks = _text_blocks(text, safe_name)
                    warnings.append(f"Text was produced by the configured OCR provider: {ocr.provider}.")
                elif ocr.safe_warning:
                    warnings.append(ocr.safe_warning)
        elif suffix == ".docx":
            import docx

            document = docx.Document(BytesIO(data))
            parts: list[str] = []
            section: str | None = None
            for paragraph in document.paragraphs:
                value = paragraph.text.strip()
                if not value:
                    continue
                style = str(getattr(paragraph.style, "name", "") or "")
                if style.lower().startswith("heading"):
                    section = value
                parts.append(value)
                blocks.append(
                    {
                        "text": value,
                        "kind": "heading" if style.lower().startswith("heading") else "paragraph",
                        "section": section,
                        "filename": safe_name,
                    }
                )
            for table_number, table in enumerate(document.tables, start=1):
                table_lines = [" | ".join(cell.text for cell in row.cells) for row in table.rows]
                value = "\n".join(table_lines)
                parts.extend(table_lines)
                blocks.append(
                    {"text": value, "kind": "table", "section": section, "table": table_number, "filename": safe_name}
                )
            text = "\n".join(parts)
        elif suffix == ".rtf":
            from striprtf.striprtf import rtf_to_text

            text = rtf_to_text(_decode_text(data))
            blocks = _text_blocks(text, safe_name)
        elif suffix in {".csv", ".tsv"}:
            import pandas as pd

            if len(data) > config.MAX_UPLOAD_BYTES:
                raise ValueError("Tabular upload exceeds the configured size limit.")
            frame = pd.read_csv(BytesIO(data), sep="\t" if suffix == ".tsv" else ",")
            _validate_frame(frame)
            text, table_blocks = _frame_to_artifact(frame, safe_name, safe_name)
            blocks.extend(table_blocks)
        elif suffix in {".xlsx", ".xls"}:
            import pandas as pd

            sheets = pd.read_excel(BytesIO(data), sheet_name=None)
            parts = []
            for sheet, frame in sheets.items():
                _validate_frame(frame)
                sheet_text, sheet_blocks = _frame_to_artifact(frame, str(sheet), safe_name)
                parts.append(sheet_text)
                blocks.extend(sheet_blocks)
            text = "\n\n".join(parts)
        elif suffix in {".json", ".jsonl"}:
            raw = _decode_text(data)
            if suffix == ".jsonl":
                values = [json.loads(line) for line in raw.splitlines() if line.strip()]
                if len(values) > config.MAX_DATASET_ROWS:
                    raise ValueError("JSONL row count exceeds the configured limit.")
            else:
                values = json.loads(raw)
            text = json.dumps(values, indent=2, ensure_ascii=False)
            blocks = _text_blocks(text, safe_name)
        elif suffix in {".html", ".htm"}:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(data, "html.parser")
            for element in soup(["script", "style", "iframe", "object", "embed", "form"]):
                element.decompose()
            text = soup.get_text("\n", strip=True)
            blocks = _html_blocks(soup, safe_name)
        else:  # .pptx
            from pptx import Presentation

            presentation = Presentation(BytesIO(data))
            slides = []
            for index, slide in enumerate(presentation.slides, start=1):
                parts = [f"Slide {index}"]
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        parts.append(shape.text.strip())
                    if getattr(shape, "has_table", False):
                        parts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                try:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        parts.append(f"Speaker notes: {notes}")
                except (AttributeError, KeyError):
                    warnings.append(f"Slide {index} speaker notes could not be extracted.")
                slide_text = "\n".join(parts)
                slides.append(slide_text)
                blocks.append(
                    {
                        "text": slide_text,
                        "kind": "slide",
                        "page": index,
                        "section": f"Slide {index}",
                        "filename": safe_name,
                    }
                )
            text = "\n\n".join(slides)
    except (zipfile.BadZipFile, json.JSONDecodeError, UnicodeError) as exc:
        raise ValueError(f"Malformed {suffix or 'uploaded'} document: {type(exc).__name__}") from exc

    text = text.strip()
    if not text:
        raise ValueError(
            f"No readable text was found in {safe_name}. Scanned PDFs and image-only files require the configured OCR extension."
        )
    if len(text) > config.MAX_EXTRACTED_CHARACTERS:
        raise ValueError(f"Extracted document exceeds the {config.MAX_EXTRACTED_CHARACTERS} character limit.")
    return {
        "filename": safe_name,
        "original_filename": filename,
        "text": text,
        "blocks": blocks or _text_blocks(text, safe_name),
        "warnings": warnings,
        "content_hash": text_hash(text),
        "partially_extracted": bool(warnings),
        "malware_scan": scan_metadata,
    }


def detect_duplicate_documents(
    documents: list[dict[str, Any]], *, near_duplicate_threshold: float = 0.92
) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    for index, document in enumerate(documents):
        for previous_index in range(index):
            previous = documents[previous_index]
            if document.get("content_hash") == previous.get("content_hash"):
                findings.append(
                    {
                        "document": document["filename"],
                        "duplicate_of": previous["filename"],
                        "type": "exact",
                        "similarity": 1.0,
                    }
                )
                break
            similarity = _token_similarity(str(document.get("text", "")), str(previous.get("text", "")))
            if similarity >= near_duplicate_threshold:
                findings.append(
                    {
                        "document": document["filename"],
                        "duplicate_of": previous["filename"],
                        "type": "near",
                        "similarity": round(similarity, 3),
                    }
                )
                break
    return findings


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _frame_to_artifact(frame, label: str, filename: str) -> tuple[str, list[dict[str, Any]]]:
    frame = frame.fillna("")
    text = f"Table: {label}\n{frame.to_csv(index=False)}"
    blocks = []
    columns = max(1, len(frame.columns))
    for start in range(0, len(frame), 100):
        stop = min(len(frame), start + 100)
        end_column = _excel_column(columns)
        blocks.append(
            {
                "text": frame.iloc[start:stop].to_csv(index=False),
                "kind": "table",
                "section": str(label),
                "sheet": str(label),
                "cell_range": f"A{start + 2}:{end_column}{stop + 1}",
                "filename": filename,
            }
        )
    return text, blocks


def _validate_frame(frame) -> None:
    if len(frame) > config.MAX_DATASET_ROWS:
        raise ValueError(f"Table has {len(frame)} rows; the configured limit is {config.MAX_DATASET_ROWS}.")
    if len(frame.columns) > 500:
        raise ValueError("Table exceeds the 500-column resource limit.")


def _validate_zip_container(data: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 5000:
                raise ValueError("Office archive contains too many entries.")
            total = sum(info.file_size for info in infos)
            compressed = max(1, sum(info.compress_size for info in infos))
            if total > config.MAX_EXTRACTED_CHARACTERS * 10 or total / compressed > 200:
                raise ValueError("Office archive exceeds decompression safety limits.")
            if any(".." in Path(info.filename).parts or Path(info.filename).is_absolute() for info in infos):
                raise ValueError("Office archive contains an unsafe path.")
    except zipfile.BadZipFile as exc:
        raise ValueError("Malformed Office archive.") from exc


def _text_blocks(text: str, filename: str, *, page: int | None = None) -> list[dict[str, Any]]:
    blocks = []
    section: str | None = None
    for part in re.split(r"\n\s*\n", text or ""):
        value = part.strip()
        if not value:
            continue
        first = value.splitlines()[0]
        if first.lstrip().startswith("#"):
            section = first.lstrip("# ")
        blocks.append({"text": value, "kind": "paragraph", "page": page, "section": section, "filename": filename})
    return blocks


def _html_blocks(soup, filename: str) -> list[dict[str, Any]]:
    blocks = []
    section = None
    for element in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "table"]):
        value = element.get_text(" ", strip=True)
        if not value:
            continue
        if element.name.startswith("h"):
            section = value
        blocks.append(
            {
                "text": value,
                "kind": "heading" if element.name.startswith("h") else element.name,
                "section": section,
                "filename": filename,
            }
        )
    return blocks


def _excel_column(number: int) -> str:
    output = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        output = chr(65 + remainder) + output
    return output or "A"


def _token_similarity(left: str, right: str) -> float:
    left_tokens = set(re.findall(r"\w+", left.lower()))
    right_tokens = set(re.findall(r"\w+", right.lower()))
    return len(left_tokens & right_tokens) / max(1, len(left_tokens | right_tokens))
