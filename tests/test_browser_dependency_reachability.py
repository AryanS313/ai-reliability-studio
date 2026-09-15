"""Portable reachability checks; native results do not certify a Wasm build.

Run this file again inside the assembled Pyodide environment. Fixtures are built
before tripwires: their creation deliberately uses image APIs the product must
not invoke when extracting document text.
"""

from __future__ import annotations

import importlib
import json
import pickle
import socket
import ssl
import urllib.request
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import numpy as np
import pandas as pd
import pytest

from src import config
from src.document_loader import extract_document_artifact
from src.domain import ExecutionStatus
from src.embeddings import LocalTfidfEmbedder
from src.reporting import json_report
from src.targets import ExternalHTTPTarget, ExternalTargetConfig, SecretResolver, TargetConfigurationError
from src.vector_store import SimpleVectorStore

POLICY = "Keep support records for 30 days."


def image_document_fixtures() -> dict[str, bytes]:
    """Normal text with embedded images, including an image XObject in the PDF."""
    import docx
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

    image = BytesIO()
    Image.new("RGB", (2, 2), color=(255, 0, 0)).save(image, format="PNG")
    document = docx.Document()
    document.add_paragraph(POLICY)
    document.add_picture(BytesIO(image.getvalue()))
    word = BytesIO()
    document.save(word)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = POLICY
    slide.shapes.add_picture(BytesIO(image.getvalue()), Inches(1), Inches(3))
    powerpoint = BytesIO()
    presentation.save(powerpoint)

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    bitmap = DecodedStreamObject()
    bitmap.set_data(b"\xff\x00\x00")
    bitmap.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(1),
            NameObject("/Height"): NumberObject(1),
            NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
            NameObject("/BitsPerComponent"): NumberObject(8),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
            NameObject("/XObject"): DictionaryObject({NameObject("/Im1"): bitmap}),
        }
    )
    content = DecodedStreamObject()
    content.set_data(f"q 20 0 0 20 10 10 cm /Im1 Do Q BT /F1 12 Tf 20 260 Td ({POLICY}) Tj ET".encode())
    page[NameObject("/Contents")] = content
    pdf = BytesIO()
    writer.write(pdf)
    return {"policy.pdf": pdf.getvalue(), "policy.docx": word.getvalue(), "policy.pptx": powerpoint.getvalue()}


def install_image_tripwires(monkeypatch) -> list[str]:
    """Fail on parsing, rendering or transforming image/font input in extraction."""
    from PIL import Image, ImageFont

    unavailable = []

    def forbidden(*args, **kwargs):
        raise AssertionError("Document text extraction reached an image/font operation")

    for owner, names in (
        (Image, ("open", "new", "frombytes", "frombuffer", "alpha_composite")),
        (Image.Image, ("save", "load", "crop", "paste", "filter")),
        (ImageFont, ("truetype", "load", "load_path", "load_default")),
    ):
        for name in names:
            if hasattr(owner, name):
                monkeypatch.setattr(owner, name, forbidden)
    try:
        cms = importlib.import_module("PIL.ImageCms")
    except ImportError:
        unavailable.append("PIL.ImageCms")
    else:
        for name in (
            "getOpenProfile",
            "buildTransform",
            "buildTransformFromOpenProfiles",
            "applyTransform",
            "profileToProfile",
        ):
            monkeypatch.setattr(cms, name, forbidden)
        monkeypatch.setattr(cms.ImageCmsTransform, "apply", forbidden)
        monkeypatch.setattr(cms.ImageCmsTransform, "apply_in_place", forbidden)
    return unavailable


@pytest.mark.parametrize("filename", ["policy.pdf", "policy.docx", "policy.pptx"])
def test_text_extraction_skips_embedded_image_and_font_operations(monkeypatch, filename):
    documents = image_document_fixtures()
    monkeypatch.setattr(config, "REQUIRE_MALWARE_SCAN", False)
    install_image_tripwires(monkeypatch)
    artifact = extract_document_artifact(filename, documents[filename])
    assert POLICY in artifact["text"]
    assert artifact["blocks"]


def entity_document(filename: str, original: bytes, destination: str) -> bytes:
    part = "word/document.xml" if filename.endswith(".docx") else "ppt/slides/slide1.xml"
    prefix = "w" if filename.endswith(".docx") else "a"
    output = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            value = source.read(name)
            if name == part:
                declaration_end = value.index(b"?>") + 2
                value = (
                    value[:declaration_end]
                    + f'<!DOCTYPE probe [<!ENTITY probe SYSTEM "{destination}">]>'.encode()
                    + value[declaration_end:]
                )
                value = value.replace(
                    f"<{prefix}:t>{POLICY}</{prefix}:t>".encode(),
                    f"<{prefix}:t>Keep source text. &probe;</{prefix}:t>".encode(),
                )
                assert b"&probe;" in value
            target.writestr(name, value)
    return output.getvalue()


@pytest.mark.parametrize("filename", ["policy.docx", "policy.pptx"])
@pytest.mark.parametrize("scheme", ["file", "https"])
def test_office_external_entities_neither_fetch_nor_substitute_secrets(tmp_path, monkeypatch, filename, scheme):
    from lxml import etree

    secret = "ENTITY_SENTINEL_6ad9e76b"
    marker = tmp_path / "not-for-document.txt"
    marker.write_text(secret)
    destination = marker.as_uri() if scheme == "file" else "https://entity.example.test/secret"
    payload = entity_document(filename, image_document_fixtures()[filename], destination)
    monkeypatch.setattr(config, "REQUIRE_MALWARE_SCAN", False)

    class RefuseResolver(etree.Resolver):
        def resolve(self, url, public_id, context):
            raise AssertionError("An office parser tried to resolve an external entity")

    modules = ("docx.oxml.parser", "docx.opc.oxml") if filename.endswith(".docx") else ("pptx.oxml",)
    resolvers = []
    for name in modules:
        parser = importlib.import_module(name).oxml_parser
        resolver = RefuseResolver()
        parser.resolvers.add(resolver)
        resolvers.append((parser, resolver))

    def no_network(*args, **kwargs):
        raise AssertionError("Document extraction must not fetch external resources")

    monkeypatch.setattr(urllib.request, "urlopen", no_network)
    monkeypatch.setattr(socket, "getaddrinfo", no_network)
    try:
        result = extract_document_artifact(filename, payload)
        assert secret not in json.dumps(result)
        assert "Keep source text." in result["text"]
    finally:
        for parser, resolver in resolvers:
            parser.resolvers.remove(resolver)
    assert marker.read_text() == secret


def test_xml_upload_is_literal_text_without_entity_expansion(tmp_path, monkeypatch):
    secret = "XML_SENTINEL_9e228"
    path = tmp_path / "secret.txt"
    path.write_text(secret)
    xml = f'<!DOCTYPE probe [<!ENTITY probe SYSTEM "{path.as_uri()}">]><policy>&probe;</policy>'
    monkeypatch.setattr(config, "REQUIRE_MALWARE_SCAN", False)
    result = extract_document_artifact("source.xml", xml.encode())
    assert result["text"] == xml
    assert secret not in result["text"]


@pytest.mark.parametrize("legacy_retained_attribute", [False, True])
def test_tfidf_discards_introspection_tokens_without_changing_vectors_or_retrieval(
    monkeypatch, legacy_retained_attribute
):
    from sklearn.feature_extraction.text import TfidfVectorizer

    chunks = [
        {"chunk_id": "records", "source_name": "Policy", "chunk_text": POLICY},
        {"chunk_id": "security", "source_name": "Policy", "chunk_text": "Escalate exposed credentials to security."},
    ]
    texts = [chunk["chunk_text"] for chunk in chunks]
    original = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True)
    expected = original.fit_transform(texts)
    embedder = LocalTfidfEmbedder()
    assert embedder.vectorizer is not None
    if legacy_retained_attribute:
        fit = embedder.vectorizer.fit_transform

        def legacy_fit(values):
            matrix = fit(values)
            embedder.vectorizer.stop_words_ = {"discarded_introspection_marker"}
            return matrix

        monkeypatch.setattr(embedder.vectorizer, "fit_transform", legacy_fit)
    actual = embedder.fit_transform(texts)
    np.testing.assert_allclose(actual.toarray(), expected.toarray())
    np.testing.assert_allclose(
        embedder.transform(["records retention"]).toarray(), original.transform(["records retention"]).toarray()
    )
    assert not hasattr(embedder.vectorizer, "stop_words_")

    def no_pickle(*args, **kwargs):
        raise AssertionError("Retrieval and evidence export must not serialize fitted vectorizers")

    monkeypatch.setattr(pickle, "dumps", no_pickle)
    store = SimpleVectorStore(embedder)
    store.build(chunks)
    found = store.retrieve("How long are support records retained?", similarity_threshold=0)
    assert found[0]["chunk_id"] == "records"
    encoded = json.dumps(found, allow_nan=False)
    assert "stop_words_" not in encoded and "discarded_introspection_marker" not in encoded
    report = json_report(
        pd.DataFrame([{"case_id": "fixture", "execution_status": "failed", "retrieved_chunks": found}])
    )
    assert "discarded_introspection_marker" not in report and "TfidfVectorizer" not in report


def test_browser_cannot_reach_native_http_key_or_certificate_operations(monkeypatch):
    from cryptography import x509
    from cryptography.hazmat.primitives import serialization

    monkeypatch.setattr(config, "APP_ACCESS_MODE", "browser")
    monkeypatch.setattr(config.sys, "platform", "emscripten")
    monkeypatch.setattr(config, "EXTERNAL_TARGET_ALLOWED_HOSTS", ("assistant.example.test",))
    monkeypatch.setenv("BROWSER_HOST_SECRET", "must-not-be-inherited")

    def forbidden(*args, **kwargs):
        raise AssertionError("Browser connection reached native network/key/certificate processing")

    for owner, names in (
        (socket, ("socket", "getaddrinfo")),
        (ssl, ("create_default_context",)),
        (urllib.request, ("urlopen",)),
        (serialization, ("load_pem_public_key", "load_der_public_key", "load_pem_private_key", "load_der_private_key")),
        (x509, ("load_pem_x509_certificate", "load_der_x509_certificate")),
    ):
        for name in names:
            monkeypatch.setattr(owner, name, forbidden)
    target = ExternalHTTPTarget(
        ExternalTargetConfig(name="No browser HTTP", endpoint="https://assistant.example.test"), opener=forbidden
    )
    assert target.execute({"question": "Fixture"}).status == ExecutionStatus.INVALID_RESPONSE
    assert target.health_check().status == ExecutionStatus.INVALID_RESPONSE
    with pytest.raises(TargetConfigurationError, match="not configured"):
        SecretResolver(environment_names=["BROWSER_HOST_SECRET"]).resolve("secret://BROWSER_HOST_SECRET")
