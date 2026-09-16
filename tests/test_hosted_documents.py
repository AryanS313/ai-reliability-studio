from __future__ import annotations

import subprocess
import zipfile
from io import BytesIO

import pytest

from src import config, hosted_documents
from src.document_loader import extract_document_artifact


@pytest.fixture(autouse=True)
def hosted_mode(monkeypatch):
    monkeypatch.setattr(config, "APP_ACCESS_MODE", "hosted-session")
    monkeypatch.setattr(config, "REQUIRE_MALWARE_SCAN", False)


@pytest.mark.parametrize(
    ("name", "data", "expected"),
    [
        ("policy.txt", b"Refunds need approval.", "Refunds need approval."),
        ("policy.csv", b"action,owner\nrefund,support\n", "refund,support"),
        ("policy.html", b"<script>alert(1)</script><h1>Refund policy</h1>", "Refund policy"),
        ("policy.json", b'{"action":"review"}', '"action": "review"'),
    ],
)
def test_real_hosted_worker_extracts_supported_text(name, data, expected):
    result = extract_document_artifact(name, data)
    assert expected in result["text"]
    assert result["filename"] == name
    assert result["blocks"]
    assert result["malware_scan"]["verdict"] == "not_requested"
    assert "alert(1)" not in result["text"]


def test_real_hosted_worker_extracts_office_documents():
    import docx
    import pandas as pd
    from pptx import Presentation

    document = docx.Document()
    document.add_paragraph("Refunds require a receipt.")
    doc_buffer = BytesIO()
    document.save(doc_buffer)
    assert "receipt" in extract_document_artifact("policy.docx", doc_buffer.getvalue())["text"]
    spreadsheet = BytesIO()
    with pd.ExcelWriter(spreadsheet, engine="openpyxl") as writer:
        pd.DataFrame({"policy": ["Escalate fraud"]}).to_excel(writer, index=False)
    assert "Escalate fraud" in extract_document_artifact("policy.xlsx", spreadsheet.getvalue())["text"]
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Review refund disputes"
    deck = BytesIO()
    presentation.save(deck)
    assert "Review refund disputes" in extract_document_artifact("policy.pptx", deck.getvalue())["text"]


def test_real_hosted_worker_extracts_pdf_text_without_images_or_ocr():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 10 100 Td (Refund window is 30 days.) Tj ET")
    page[NameObject("/Contents")] = content
    buffer = BytesIO()
    writer.write(buffer)
    artifact = extract_document_artifact("policy.pdf", buffer.getvalue())
    assert "Refund window is 30 days." in artifact["text"]
    assert artifact["blocks"][0]["page"] == 1


def test_hosted_office_decompression_bomb_rejected_before_office_parser():
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 3_000_000)
    with pytest.raises(ValueError, match="decompression safety limits"):
        extract_document_artifact("large.docx", buffer.getvalue())


@pytest.mark.parametrize("format_name", ["csv", "tsv", "xlsx", "json", "jsonl"])
def test_hosted_question_import_keeps_real_case_values(format_name):
    import pandas as pd

    frame = pd.DataFrame(
        [
            {
                "question": "Can I get a refund?",
                "expected_answer": "Within 30 days.",
                "expected_source": "Refund policy",
                "category": "refund",
                "should_escalate": False,
            }
        ]
    )
    if format_name in {"csv", "tsv"}:
        data = frame.to_csv(index=False, sep="\t" if format_name == "tsv" else ",").encode()
    elif format_name == "xlsx":
        buffer = BytesIO()
        frame.to_excel(buffer, index=False)
        data = buffer.getvalue()
    else:
        data = frame.to_json(orient="records", lines=format_name == "jsonl").encode()
    result = hosted_documents.read_hosted_dataset(f"questions.{format_name}", data)
    assert result.iloc[0]["question"] == "Can I get a refund?"
    assert not result.iloc[0]["should_escalate"]
    assert result.iloc[0]["case_id"]


def test_hosted_question_import_limits_rows_before_loading_whole_table():
    data = b"question,expected_answer,expected_source,category,should_escalate\n" + b"q,a,s,c,false\n" * 501
    with pytest.raises(ValueError, match="at most 500 questions"):
        hosted_documents.read_hosted_dataset("questions.csv", data)


@pytest.mark.parametrize(
    ("name", "data", "message"),
    [
        ("empty.txt", b"", "Empty"),
        ("bad.exe", b"MZ", "Unsupported"),
        ("bad.docx", b"not a ZIP archive", "Malformed"),
        ("bad.json", b"{", "Malformed"),
        ("old.doc", b"legacy", "Save the file as .docx"),
    ],
)
def test_hosted_upload_rejects_empty_unsupported_and_malformed(name, data, message):
    with pytest.raises(ValueError, match=message):
        extract_document_artifact(name, data)


def test_hosted_extracted_size_cannot_exceed_parent_limit(monkeypatch):
    monkeypatch.setattr(config, "MAX_EXTRACTED_CHARACTERS", 20)
    with pytest.raises(ValueError, match="character limit"):
        extract_document_artifact("policy.txt", b"x" * 21)


def test_hosted_hard_file_limit_cannot_be_relaxed(monkeypatch):
    monkeypatch.setattr(config, "MAX_UPLOAD_BYTES", 20 * 1024 * 1024)
    with pytest.raises(ValueError, match="2 MB"):
        extract_document_artifact("policy.txt", b"x" * (2 * 1024 * 1024 + 1))


def test_required_scanner_is_enforced_before_hosted_worker(monkeypatch):
    monkeypatch.setattr(config, "REQUIRE_MALWARE_SCAN", True)
    monkeypatch.setattr(hosted_documents.subprocess, "run", lambda *a, **kw: pytest.fail("must not parse"))
    with pytest.raises(ValueError, match="requires a configured malware scanner"):
        extract_document_artifact("policy.txt", b"Sample policy")


@pytest.mark.parametrize("filename", ["questions.csv", "questions.xlsx"])
def test_required_scanner_blocks_question_files_before_parser(monkeypatch, filename):
    monkeypatch.setattr(config, "REQUIRE_MALWARE_SCAN", True)
    monkeypatch.setattr(hosted_documents.subprocess, "run", lambda *a, **kw: pytest.fail("must not parse"))
    with pytest.raises(ValueError, match="requires a configured malware scanner"):
        hosted_documents.read_hosted_dataset(filename, b"unscanned data")


def test_parser_process_does_not_inherit_keys_and_recovers_after_timeout(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "owner-secret-must-not-inherit")
    monkeypatch.setenv("HTTPS_PROXY", "https://private-proxy.invalid")
    original = hosted_documents.subprocess.run

    def timeout(command, **kwargs):
        assert "OPENAI_API_KEY" not in kwargs["env"]
        assert "HTTPS_PROXY" not in kwargs["env"]
        assert kwargs["env"]["STUDIO_EXTRACTION_WORKER"] == "1"
        assert command[1] == "-I"
        assert "Sample policy" not in str(command)
        raise subprocess.TimeoutExpired(command, kwargs["timeout"])

    monkeypatch.setattr(hosted_documents.subprocess, "run", timeout)
    with pytest.raises(ValueError, match="took too long"):
        extract_document_artifact("policy.txt", b"Sample policy")
    monkeypatch.setattr(hosted_documents.subprocess, "run", original)
    assert extract_document_artifact("policy.txt", b"Recovered")["text"] == "Recovered"


def test_saturated_parser_slots_reject_instead_of_queueing_unbounded_work():
    assert hosted_documents._slots.acquire(blocking=False)
    try:
        with pytest.raises(ValueError, match="processing is busy"):
            extract_document_artifact("policy.txt", b"Sample policy")
    finally:
        hosted_documents._slots.release()
