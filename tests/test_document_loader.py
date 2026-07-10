from io import BytesIO

import pandas as pd

from src.document_loader import extract_document_bytes


def test_extracts_plain_text_csv_json_and_html():
    assert extract_document_bytes("policy.txt", b"Refunds require review.") == "Refunds require review."
    assert "refund" in extract_document_bytes("cases.csv", b"topic,action\nrefund,review\n")
    assert '"policy"' in extract_document_bytes("policy.json", b'{"policy": "KYC"}')
    assert extract_document_bytes("policy.html", b"<h1>KYC</h1><p>Manual review</p>") == "KYC\nManual review"


def test_extracts_docx_tables():
    import docx

    document = docx.Document()
    document.add_paragraph("Account closure policy")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Blocker"
    table.cell(0, 1).text = "Active loan"
    buffer = BytesIO()
    document.save(buffer)

    text = extract_document_bytes("policy.docx", buffer.getvalue())
    assert "Account closure policy" in text
    assert "Blocker | Active loan" in text


def test_extracts_all_excel_sheets():
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame({"rule": ["Escalate fraud"]}).to_excel(writer, sheet_name="Escalation", index=False)
        pd.DataFrame({"rule": ["Verify PAN"]}).to_excel(writer, sheet_name="KYC", index=False)

    text = extract_document_bytes("policies.xlsx", buffer.getvalue())
    assert "Table: Escalation" in text
    assert "Escalate fraud" in text
    assert "Table: KYC" in text


def test_extracts_powerpoint_slide_text():
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Refund policy"
    slide.placeholders[1].text = "Escalate suspected fraud"
    buffer = BytesIO()
    presentation.save(buffer)

    text = extract_document_bytes("policy.pptx", buffer.getvalue())
    assert "Slide 1" in text
    assert "Refund policy" in text
    assert "Escalate suspected fraud" in text


def test_legacy_doc_has_actionable_error():
    try:
        extract_document_bytes("old-policy.doc", b"legacy")
    except ValueError as exc:
        assert "Save the file as .docx" in str(exc)
    else:
        raise AssertionError("Legacy .doc should be rejected with conversion guidance")
