"""Source review must preserve policy values, structure and source identity."""

import csv
import json
from io import BytesIO, StringIO

import pytest

from src.chunker import chunk_document
from src.document_loader import extract_document_artifact
from src.evaluator import read_eval_dataset
from src.saved_responses import read_response_file
from src.ui_workflows import (
    new_review_workspace,
    read_reference_json,
    read_reference_uploads,
    read_review_workspace,
    source_extraction_warnings,
)


def policy_docx(*, table_after_second_heading=False):
    import docx

    document = docx.Document()
    document.add_heading("Standard workspaces", level=1)
    if table_after_second_heading:
        document.add_heading("Regulated workspaces", level=1)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Audit log retention"
    table.cell(0, 1).text = "30 days"
    if not table_after_second_heading:
        document.add_heading("Regulated workspaces", level=1)
    document.add_paragraph("Legal hold overrides automatic deletion.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_table_stays_with_its_original_heading_and_reading_order():
    artifact = extract_document_artifact("retention.docx", policy_docx())
    table = next(block for block in artifact["blocks"] if block["kind"] == "table")
    assert table["section"] == "Standard workspaces"
    assert artifact["text"].index("Standard workspaces") < artifact["text"].index("Audit log retention")
    assert artifact["text"].index("Audit log retention") < artifact["text"].index("Regulated workspaces")
    table_chunk = next(chunk for chunk in chunk_document(artifact) if chunk["kind"] == "table")
    assert table_chunk["section"] == "Standard workspaces"


def test_moving_docx_policy_table_to_another_section_changes_source_identity():
    original = extract_document_artifact("retention.docx", policy_docx())
    moved = extract_document_artifact("retention.docx", policy_docx(table_after_second_heading=True))
    assert original["content_hash"] != moved["content_hash"]
    assert {chunk["document_hash"] for chunk in chunk_document(original)} != {
        chunk["document_hash"] for chunk in chunk_document(moved)
    }


@pytest.mark.parametrize(("suffix", "delimiter"), [("csv", ","), ("tsv", "\t")])
def test_source_table_preserves_literal_codes_missing_value_words_and_decimal_text(suffix, delimiter):
    rows = [["plan_id", "region", "manual_value", "limit"], ["001", "NA", "NULL", "1.00"]]
    buffer = StringIO()
    csv.writer(buffer, delimiter=delimiter).writerows(rows)
    artifact = extract_document_artifact(f"policy.{suffix}", buffer.getvalue().encode())
    extracted_rows = list(csv.reader(StringIO(artifact["blocks"][0]["text"])))
    assert extracted_rows == rows


def test_excel_source_preserves_text_cell_values():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["plan_id", "region", "limit"])
    sheet.append(["001", "NA", "1.00"])
    buffer = BytesIO()
    workbook.save(buffer)
    artifact = extract_document_artifact("policy.xlsx", buffer.getvalue())
    assert list(csv.reader(StringIO(artifact["blocks"][0]["text"]))) == [
        ["plan_id", "region", "limit"],
        ["001", "NA", "1.00"],
    ]


@pytest.mark.parametrize(
    ("filename", "text", "encoding"),
    [
        ("policy.txt", "Café", "latin-1"),
        ("policy.txt", "Café support — retain 30 days.", "utf-8-sig"),
        ("policy.md", "# Limits\n\nDo not exceed 30 days — except legal hold.", "utf-16"),
    ],
)
def test_text_decoding_preserves_supported_text_instead_of_guessing_utf16(filename, text, encoding):
    assert extract_document_artifact(filename, text.encode(encoding))["text"] == text


def text_and_blank_pdf():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 260 Td (Plan | Retention) Tj 0 -20 Td (Basic | 30 days) Tj ET")
    page[NameObject("/Contents")] = stream
    writer.add_blank_page(width=300, height=300)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_pdf_preserves_simple_table_text_page_and_explicit_unreadable_page_warning():
    artifact = extract_document_artifact("policy.pdf", text_and_blank_pdf())
    assert "Plan | Retention" in artifact["text"] and "Basic | 30 days" in artifact["text"]
    assert artifact["blocks"][0]["page"] == 1
    assert artifact["partially_extracted"]
    assert any("Page 2 contains no searchable text" in warning for warning in artifact["warnings"])


def test_structured_source_json_retains_passage_text_ids_and_declared_versions():
    passages = [
        {
            "chunk_id": "001",
            "source_name": "NA support policy",
            "document_id": "policy-01",
            "document_version": "edition-007",
            "section": "Quoted values",
            "chunk_text": "The code is NULL.\nCafé support keeps the display value 1.00.",
        }
    ]
    assert read_reference_json(json.dumps(passages).encode()) == passages


@pytest.mark.parametrize("extension", ["csv", "json"])
def test_question_and_response_case_identifiers_roundtrip_without_numeric_coercion(extension):
    rows = [
        {
            "case_id": "001",
            "question": "Which support code applies?",
            "expected_answer": "001",
            "expected_source": "NA",
            "category": "routine",
            "should_escalate": False,
        }
    ]
    if extension == "csv":
        buffer = StringIO()
        writer = csv.DictWriter(buffer, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
        content = buffer.getvalue().encode()
    else:
        content = json.dumps(rows).encode()
    questions = read_eval_dataset(f"questions.{extension}", content)
    answers = read_response_file("answers.csv", b"case_id,actual_answer\n001,The support code is 001.\n")
    assert questions.iloc[0]["case_id"] == answers[0]["case_id"] == "001"
    assert questions.iloc[0]["expected_answer"] == "001"
    assert questions.iloc[0]["expected_source"] == "NA"


def test_saved_source_import_and_workspace_restore_preserve_partial_extraction_notices():
    import pandas as pd

    upload = BytesIO(text_and_blank_pdf())
    upload.name = "policy.pdf"
    chunks = read_reference_uploads([upload])
    notices = source_extraction_warnings(chunks)
    assert notices == ["policy.pdf: Page 2 contains no searchable text; configure the OCR extension for scanned pages."]
    assert all(chunk["partially_extracted"] is True for chunk in chunks)
    assert source_extraction_warnings([*chunks, *chunks]) == notices
    dataset = pd.DataFrame(
        [
            {
                "case_id": "source-warning-case",
                "question": "How long are Basic plan records retained?",
                "expected_answer": "30 days",
                "expected_source": chunks[0]["source_name"],
                "category": "routine",
                "should_escalate": False,
            }
        ]
    )
    workspace = new_review_workspace(
        dataset,
        chunks,
        [{"case_id": "source-warning-case", "actual_answer": "30 days"}],
        target_name="Fictional source review",
        target_version="fixture-v1",
        captured_at="2026-01-01T12:00:00+00:00",
        evidence_kind="fixture",
    )
    restored = read_review_workspace(json.dumps(workspace).encode())
    assert restored["sources"] == chunks
    assert source_extraction_warnings(restored["sources"]) == notices


def test_duplicate_source_documents_do_not_discard_extraction_notices():
    uploads = []
    for name in ("first.pdf", "second.pdf"):
        upload = BytesIO(text_and_blank_pdf())
        upload.name = name
        uploads.append(upload)
    chunks = read_reference_uploads(uploads)
    assert len(chunks) == 1
    notices = source_extraction_warnings(chunks)
    assert len(notices) == 2
    assert notices[0].startswith("first.pdf:") and notices[1].startswith("second.pdf:")


def test_extraction_notice_display_is_bounded_text_and_idempotent():
    chunk = {"extraction_warnings": [None, 7, "", "X" * 900, *[f"Warning {index}" for index in range(30)]]}
    notices = source_extraction_warnings([chunk, chunk])
    assert len(notices) == 20
    assert all(isinstance(notice, str) and len(notice) <= 600 for notice in notices)
    assert notices[0].endswith("(notice shortened)")
    assert "omitted" in notices[-1]
    assert source_extraction_warnings([{"extraction_warnings": notices}]) == notices


def test_extraction_notice_helper_accepts_document_warnings_and_remains_ui_compatible():
    from src.document_loader import source_extraction_warnings as document_notices

    document = {"filename": "policy.pdf", "warnings": [None, 7, "", " Page 2 has no searchable text. "]}
    notices = document_notices([document])
    assert source_extraction_warnings is document_notices
    assert notices == ["policy.pdf: Page 2 has no searchable text."]
    assert document_notices([document, {"extraction_warnings": notices}]) == notices
    assert document["warnings"][-1] == " Page 2 has no searchable text. "


def test_fully_extracted_plain_source_does_not_gain_partial_extraction_flags():
    upload = BytesIO(b"All Basic plan records are retained for 30 days.")
    upload.name = "policy.txt"
    chunks = read_reference_uploads([upload])
    assert source_extraction_warnings(chunks) == []
    assert all("partially_extracted" not in chunk and "extraction_warnings" not in chunk for chunk in chunks)
